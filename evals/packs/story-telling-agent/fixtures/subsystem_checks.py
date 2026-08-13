"""Executable fixtures for story-telling-agent structural evals."""
import argparse, importlib.util, json, os, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[4]
PACK = REPO / "agent-packs/story-telling-agent"
SCRIPTS = PACK / "scripts"
sys.path.insert(0, str(SCRIPTS))
from common import canonical, dump, load, sha_bytes, sha_file
from validate_contract import RECEIPT_RULES, semantic_validate, schema_validate, validate_receipt_lineage

def deck(formats="both", chart=False):
    content = {"subtitle": "A sourced decision narrative"}
    typ, recipe, family, order = "title", "title", "hero", ["title", "assertion", "visual", "implication", "source-footer"]
    if chart:
        typ, recipe, family = "chart", "chart", "data"
        content = {"relationship": "distribution", "chart_type": "bar", "categories": ["A", "B"], "series": [{"name": "Count", "values": [2, 4]}], "units": "items", "source_ids": [], "evidence_ids": [], "focal_datum": "B is twice A", "annotation": "B leads", "axis_policy": "zero", "missing_value_policy": "gap", "color_semantics": "categorical", "render_mode": "native-editable", "zero_baseline_policy": "required"}
        order = ["title", "assertion", "chart", "content-annotation", "implication", "source-footer"]
    return {"schema_version": "3.0", "deck_id": "fixture", "render_seed": 7, "aspect_ratio": "16:9", "design_system_id": "executive-navy", "formats": formats, "metadata": {"title": "Fixture", "language": "en", "created_from_run_id": "run-fixture"}, "sources": [], "slides": [{"index": 1, "slide_id": "s1", "type": typ, "title": "Choose the testable path", "core_assertion": "Executable checks catch contract drift", "claim_ids": [], "evidence_ids": [], "implication": "Run the subsystem", "layout_family": family, "recipe": recipe, "background_role": "background", "content": content, "notes": {"intent": "Decide", "talk_track": "Run it", "claim_ids": [], "evidence_ids": []}, "source_footer": [], "accessibility": {"alt_text": "Decision slide", "decorative": False, "contrast_intent": "AA", "non_color_encoding": True}, "reading_order": order}]}

def state(phase="new"):
    zero = {k: 0 for k in ("intake_producer_retry", "intake_validator_contract_retry", "strategy_producer_retry", "strategy_validator_contract_retry", "render_retry", "quality_retry", "qa_contract_retry", "publication_contract_retry", "publication_retry")}
    return {"schema_version": "3.0", "run_id": "run-fixture", "created_at": "2026-01-01T00:00:00Z", "updated_at": "2026-01-01T00:00:00Z", "phase": phase, "requested_formats": "both", "effective_formats": "both", "input_manifest": [], "artifact_manifest": [], "staging_manifest": [], "publication_manifest": [], "events": [], "validation_receipt_ids": [], "validation_receipts": {}, "capability_preflight": {"status": "ok", "checked_at": "2026-01-01T00:00:00Z", "evidence_id": "cap-ok", "available_formats": "both", "capabilities": {}}, "retry_counters": zero, "invalidated_artifact_ids": [], "transition_history": [], "publication_gate": "none", "publication_destination": {"intake_output_dir": "out", "intake_sha256": "sha256:" + "1" * 64, "effective_output_dir": "out", "effective_destination_event_id": "initial-destination", "revision": 0}, "active_lineage_id": "lineage-1"}

def indexed_receipt(receipt_id, handoff_id, kind, artifact_id, checksum, operation, predecessor,
                    producer="deck-composer", validator="deck-critic", capability="not-applicable",
                    manifest=None, approval=None, acceptance=None):
    return {
        "receipt_id": receipt_id, "run_id": "run-fixture", "handoff_id": handoff_id,
        "artifact_kind": kind, "artifact_id": artifact_id, "artifact_sha256": checksum,
        "producer_agent": producer, "validator_agent": validator,
        "validator_operation": operation, "predecessor_event_id": predecessor,
        "lineage_id": "lineage-1", "valid": True, "superseded": False,
        "capability_status": capability, "render_manifest_sha256": manifest,
        "approval_event_id": approval, "acceptance_event_id": acceptance,
    }

def run(cmd, expected=0):
    result = subprocess.run(cmd, cwd=REPO, capture_output=True, text=True)
    if result.returncode != expected:
        raise AssertionError(f"{cmd}\nexit={result.returncode}\nstdout={result.stdout}\nstderr={result.stderr}")
    return result

def contract_case():
    valid = deck(chart=True)
    assert not schema_validate(valid, load(PACK / "schemas/v3/deck-spec.schema.json"))
    assert not semantic_validate("deck-spec", valid)
    invalid = json.loads(json.dumps(valid)); invalid["slides"][0]["content"]["chart_type"] = "pie"
    assert any(x["code"] == "chart-relationship-type" for x in semantic_validate("deck-spec", invalid))
    ledger = {"evidence": [{"evidence_id": "e1", "source_id": "dup"}, {"evidence_id": "e2", "source_id": "dup"}], "claims": []}
    assert any(x["code"] == "source-id-duplicate" for x in semantic_validate("evidence-ledger", ledger))

def render_case():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td); spec = td / "deck.json"; dump(spec, deck())
        run([sys.executable, str(SCRIPTS / "render_deck.py"), "--spec", str(spec), "--output-dir", str(td / "out")])
        manifest = load(td / "out/render-manifest.json")
        assert manifest["parity"]["status"] == "pass"
        assert manifest["reproducibility"]["second_render_checked"] and manifest["reproducibility"]["matching"]
        assert all(x["status"] == "rendered" for x in manifest["outputs"])

def render_variants_case():
    from pptx import Presentation
    variants = [
        ("categorical", "native-editable", "zero", "distribution", "bar", ["A", "B"], [2, 4]),
        ("sequential", "deterministic-image", "auto", "trend", "line", ["1", "2"], [2, 4]),
        ("diverging", "native-editable", "explicit", "correlation", "scatter", ["1", "3"], [2, 8]),
        ("status", "native-editable", "zero", "comparison", "column", ["A", "B"], [2, 4]),
    ]
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        for i, (colors, mode, axis, relationship, chart_type, categories, values) in enumerate(variants):
            spec = deck()
            slide = spec["slides"][0]; slide.update({"type": "chart", "recipe": "chart", "layout_family": "data", "reading_order": ["title","assertion","chart","content-annotation","implication","source-footer"]})
            content = {"relationship": relationship, "chart_type": chart_type, "categories": categories, "series": [{"name": "Series", "values": values}], "units": "score", "source_ids": [], "evidence_ids": [], "focal_datum": "fixture", "annotation": "fixture", "axis_policy": axis, "missing_value_policy": "gap", "color_semantics": colors, "render_mode": mode, "zero_baseline_policy": "required" if axis == "zero" else "not-applicable"}
            if axis == "explicit":
                content["explicit_axis"] = {"x_min": 0, "x_max": 4, "y_min": 0, "y_max": 10, "x_title": "Input", "y_title": "Outcome"}
            slide["content"] = content
            sp, pp, mp = td/f"{i}.json", td/f"{i}.pptx", td/f"{i}.md"; dump(sp, spec)
            run([sys.executable, str(SCRIPTS/"render_pptx.py"), "--spec", str(sp), "--output", str(pp)])
            run([sys.executable, str(SCRIPTS/"render_marp.py"), "--spec", str(sp), "--output", str(mp)])
            prs = Presentation(pp); chart_shape = next(x for x in prs.slides[0].shapes if x.name == "chart")
            descr = chart_shape._element.xpath(".//p:cNvPr")[0].get("descr")
            assert colors in descr and ("deterministic image" in descr) == (mode == "deterministic-image")
            assert chart_shape.has_chart == (mode == "native-editable")
            marp = mp.read_text(encoding="utf-8")
            assert f'data-color-semantics="{colors}"' in marp and f'data-render-mode="{mode}"' in marp
            if chart_type == "scatter":
                with zipfile.ZipFile(pp) as z:
                    xml = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("ppt/charts/chart"))
                assert "<c:xVal>" in xml and "Input" in xml and "Outcome" in xml
                assert 'data-axis-policy="explicit"' in marp and "X: Input [0, 4]" in marp
        spec = deck(); slide = spec["slides"][0]
        slide.update({"type":"matrix","recipe":"matrix","layout_family":"comparison","reading_order":["title","assertion","matrix","matrix-x","matrix-y","low","high","implication","source-footer"],"content":{"x_axis":"Effort","y_axis":"Impact","items":[{"object_id":"low","label":"Low","x":.1,"y":.2},{"object_id":"high","label":"High","x":.9,"y":.8}]}})
        sp, pp, mp = td/"matrix.json", td/"matrix.pptx", td/"matrix.md"; dump(sp,spec)
        run([sys.executable,str(SCRIPTS/"render_pptx.py"),"--spec",str(sp),"--output",str(pp)])
        run([sys.executable,str(SCRIPTS/"render_marp.py"),"--spec",str(sp),"--output",str(mp)])
        shapes={x.name:x for x in Presentation(pp).slides[0].shapes}
        assert shapes["low"].left < shapes["high"].left and shapes["low"].top > shapes["high"].top
        marp=mp.read_text(encoding="utf-8")
        assert 'data-x="0.1" data-y="0.2"' in marp and "left:90.00%;bottom:80.00%" in marp

def all_charts_case():
    from io import BytesIO
    from PIL import Image
    from pptx import Presentation
    allowed = {
        "comparison": ("bar", "column"),
        "trend": ("line", "area", "column"),
        "composition": ("pie", "donut", "bar"),
        "distribution": ("bar", "column"),
        "correlation": ("scatter",),
        "flow": ("bar",),
    }
    all_types = {"bar", "column", "line", "area", "pie", "donut", "scatter"}
    with tempfile.TemporaryDirectory() as td:
        td = Path(td); hashes = {}
        for relationship, chart_types in allowed.items():
            for chart_type in chart_types:
                for mode in ("native-editable", "deterministic-image"):
                    spec = deck(); slide = spec["slides"][0]
                    slide.update({"type":"chart","recipe":"chart","layout_family":"data","reading_order":["title","assertion","chart","content-annotation","implication","source-footer"]})
                    categories = ["1","2","3"] if chart_type == "scatter" else ["A","B","C"]
                    slide["content"] = {
                        "relationship":relationship, "chart_type":chart_type, "categories":categories,
                        "series":[{"name":"Series","values":[2,5,3]}], "units":"items",
                        "source_ids":[], "evidence_ids":[], "focal_datum":"B", "annotation":"B leads",
                        "axis_policy":"zero" if chart_type in ("bar","column") else "auto",
                        "missing_value_policy":"gap", "color_semantics":"categorical",
                        "render_mode":mode, "zero_baseline_policy":"required" if relationship in ("comparison","distribution") else "not-applicable",
                    }
                    assert not schema_validate(spec, load(PACK/"schemas/v3/deck-spec.schema.json"))
                    assert not semantic_validate("deck-spec", spec)
                    stem=f"{relationship}-{chart_type}-{mode}"; sp=td/(stem+".json"); pp=td/(stem+".pptx"); dump(sp,spec)
                    run([sys.executable,str(SCRIPTS/"render_pptx.py"),"--spec",str(sp),"--output",str(pp)])
                    shape=next(x for x in Presentation(pp).slides[0].shapes if x.name=="chart")
                    assert shape.has_chart == (mode=="native-editable")
                    if mode=="deterministic-image":
                        image=Image.open(BytesIO(shape.image.blob)).convert("RGB")
                        colors=image.getcolors(image.width*image.height) or []
                        assert len(colors) >= 3 and max(count for count,_ in colors) < image.width*image.height
                        hashes[(relationship,chart_type)] = sha_bytes(shape.image.blob)
                        center=image.getpixel((image.width//2,image.height//2))
                        if chart_type=="donut":
                            assert center == tuple(image.getpixel((0,0)))
                        if chart_type=="pie":
                            assert center != tuple(image.getpixel((0,0)))
        assert hashes[("trend","line")] != hashes[("trend","area")]
        assert hashes[("comparison","bar")] != hashes[("comparison","column")]
        assert hashes[("composition","pie")] != hashes[("composition","donut")]
        for relationship, valid_types in allowed.items():
            for invalid_type in sorted(all_types-set(valid_types)):
                spec=deck(chart=True); spec["slides"][0]["content"].update({"relationship":relationship,"chart_type":invalid_type})
                assert any(x["code"]=="chart-relationship-type" for x in semantic_validate("deck-spec",spec)), (relationship,invalid_type)

def all_recipes_case():
    recipes = {
        "title": ({"subtitle":"Subtitle"}, "hero", ["visual"]),
        "section": ({"section_label":"Section"}, "hero", ["visual"]),
        "assertion-evidence": ({"evidence_blocks":[{"object_id":"e1","text":"Evidence"}]}, "editorial", ["e1"]),
        "big-number": ({"value":"42","unit":"%","context":"Context"}, "data", ["visual","content-1"]),
        "comparison": ({"left":"A","right":"B","comparison_basis":"Basis"}, "comparison", ["content-1","content-2","comparison-basis"]),
        "process": ({"steps":[{"object_id":"p1","label":"One"},{"object_id":"p2","label":"Two"}]}, "process", ["p1","p2"]),
        "timeline": ({"events":[{"object_id":"t1","date":"Q1","label":"One"},{"object_id":"t2","date":"Q2","label":"Two"}]}, "process", ["t1","t2"]),
        "chart": ({"relationship":"distribution","chart_type":"bar","categories":["A","B"],"series":[{"name":"Series","values":[1,2]}],"units":"items","source_ids":[],"evidence_ids":[],"focal_datum":"B","annotation":"B leads","axis_policy":"zero","missing_value_policy":"gap","color_semantics":"categorical","render_mode":"native-editable","zero_baseline_policy":"required"}, "data", ["chart","content-annotation"]),
        "image-led": ({"source":"missing-local-fixture.png","license":"fixture","crop_focal_point":"50,50","alt_text":"Fixture image","decorative":False,"fallback":"No image"}, "visual", ["image","content-1"]),
        "quote": ({"quote":"A useful quote","attribution":"Fixture"}, "editorial", ["quote","content-attribution"]),
        "matrix": ({"x_axis":"Effort","y_axis":"Impact","items":[{"object_id":"m1","label":"Choice","x":.5,"y":.5}]}, "comparison", ["matrix","matrix-x","matrix-y","m1"]),
        "table": ({"headers":[{"object_id":"h1","label":"Header","semantic_role":"header"}],"rows":[[{"object_id":"c1","value":"Value","semantic_role":"data"}]]}, "data", ["table"]),
        "recommendation": ({"recommendation":"Proceed","rationale":"Evidence","owner":"Owner","next_step":"Start"}, "decision", ["visual","content-1","content-2"]),
        "cta": ({"cta":"Approve","owner":"Owner","timing":"Today"}, "decision", ["visual","content-1","content-2"]),
    }
    with tempfile.TemporaryDirectory() as td:
        td=Path(td)
        for index,(typ,(content,family,objects)) in enumerate(recipes.items()):
            spec=deck(); slide=spec["slides"][0]
            slide.update({"type":typ,"recipe":typ,"layout_family":family,"content":content,"reading_order":["title","assertion",*objects,"implication","source-footer"]})
            sp=td/f"{index}-{typ}.json"; dump(sp,spec)
            assert not schema_validate(spec,load(PACK/"schemas/v3/deck-spec.schema.json")), typ
            assert not semantic_validate("deck-spec",spec), typ
            run([sys.executable,str(SCRIPTS/"render_deck.py"),"--spec",str(sp),"--output-dir",str(td/f"out-{index}")])

def receipt_case():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td); spec = td / "deck.json"; dump(spec, deck("pptx"))
        receipt = td / "receipt.json"
        run([sys.executable, str(SCRIPTS / "validate_contract.py"), "--kind", "deck-spec", "--input", str(spec), "--schema", str(PACK / "schemas/v3/deck-spec.schema.json"), "--run-id", "run-fixture", "--handoff-id", "h1", "--producer", "deck-composer", "--validator", "deck-critic", "--operation", "review", "--receipt", str(receipt)])
        assert load(receipt)["valid"]
        bad = load(spec); bad["slides"][0]["reading_order"] = ["title"]; dump(spec, bad)
        bad_receipt = td / "bad.json"
        run([sys.executable, str(SCRIPTS / "validate_contract.py"), "--kind", "deck-spec", "--input", str(spec), "--schema", str(PACK / "schemas/v3/deck-spec.schema.json"), "--run-id", "run-fixture", "--handoff-id", "h2", "--producer", "deck-composer", "--validator", "deck-critic", "--operation", "review", "--receipt", str(bad_receipt)])
        assert not load(bad_receipt)["valid"] and load(bad_receipt)["diagnostic_count"] > 0

def receipt_diagnostics_case():
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as td:
        td=Path(td); spec=td/"deck.json"; dump(spec,deck())
        run([sys.executable,str(SCRIPTS/"render_deck.py"),"--spec",str(spec),"--output-dir",str(td/"out")])
        manifest_path=td/"out/render-manifest.json"; manifest=load(manifest_path)
        pptx_entry=next(x for x in manifest["outputs"] if x["channel"]=="pptx")
        prs=Presentation(pptx_entry["path"]); prs.slides[0].shapes[0].name="corrupt-reading-order"; prs.save(pptx_entry["path"])
        pptx_entry["sha256"]=sha_file(pptx_entry["path"]); dump(manifest_path,manifest)
        model=td/"model.json"; dump(model,{"status":"completed","scores":{"narrative":1,"visual_craft":1,"evidence_fit":1},"notes":["fixture"]})
        report=td/"qa.json"; receipts=td/"receipts"
        result=subprocess.run([sys.executable,str(SCRIPTS/"inspect_deck.py"),"--spec",str(spec),"--manifest",str(manifest_path),"--output",str(report),"--run-id","run-fixture","--handoff-id","qa-invalid","--receipt-dir",str(receipts),"--model-review",str(model)],cwd=REPO,capture_output=True,text=True)
        assert result.returncode==1, result.stdout+result.stderr
        staged=next(load(x) for x in receipts.glob("*.json") if load(x)["artifact_kind"]=="staged-pptx")
        assert not staged["valid"] and staged["diagnostic_count"] > 0
        assert staged["schema_path"] is None
        assert staged["inspection_profile"]["id"] == "story-staged-pptx/1"
        profile = staged["inspection_profile"]
        if profile["capability_status"] == "complete":
            assert profile["completed_diagnostics"] == ["identity","structural","semantic","raster-150dpi","parity","reproducibility"]
            assert profile["unavailable_diagnostics"] == []
        else:
            assert profile["capability_status"] == "raster-unavailable"
            assert "raster-150dpi" not in profile["completed_diagnostics"]
            assert profile["unavailable_diagnostics"][0]["diagnostic"] == "raster-150dpi"
            assert not staged["valid"]

def receipt_gate_matrix_case():
    """Every receipt-gated transition accepts only its complete exact handoff."""
    for gate, rule in RECEIPT_RULES.items():
        fixture = state()
        predecessor_type = sorted(rule["predecessors"])[0]
        predecessor_id = f"before-{gate}"
        fixture["events"] = [{
            "event_id": predecessor_id, "type": predecessor_type, "run_id": fixture["run_id"],
            "at": "2026-01-01T00:00:00Z", "actor": "fixture", "evidence_ids": ["fixture"],
        }]
        manifest_name = next(iter(rule["manifests"]))
        kinds = sorted(rule["required_kinds"])
        for index, kind in enumerate(kinds):
            artifact_id = f"{gate}-{kind}"
            checksum = "sha256:" + format(index + 1, "x") * 64
            fixture[manifest_name].append({
                "artifact_id": artifact_id, "kind": kind, "path": artifact_id,
                "sha256": checksum, "producer": rule["producer"],
                "status": "staged" if manifest_name == "staging_manifest" else "validated",
                "lineage_id": fixture["active_lineage_id"],
            })
            receipt_id = f"receipt-{artifact_id}"
            fixture["validation_receipt_ids"].append(receipt_id)
            fixture["validation_receipts"][receipt_id] = indexed_receipt(
                receipt_id, f"handoff-{gate}", kind, artifact_id, checksum,
                sorted(rule["operations"])[0], predecessor_id,
                producer=rule["producer"], validator=rule["validator"],
                capability=rule.get("capability", "not-applicable"),
                manifest="sha256:" + "b" * 64 if "render-manifest" in rule["kinds"] else None,
            )
        event = {
            "type": gate, "run_id": fixture["run_id"], "handoff_id": f"handoff-{gate}",
            "receipt_ids": list(fixture["validation_receipt_ids"]),
        }
        assert validate_receipt_lineage(fixture, event) is None, gate

        def rejected(mutator):
            broken = json.loads(json.dumps(fixture))
            candidate = json.loads(json.dumps(event))
            mutator(broken, candidate)
            assert validate_receipt_lineage(broken, candidate), (gate, broken, candidate)

        rejected(lambda s, e: e["receipt_ids"].pop())  # missing
        rejected(lambda s, e: e["receipt_ids"].append("unrelated"))  # extra/missing envelope
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(lineage_id="stale"))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(artifact_kind="wrong-kind"))
        rejected(lambda s, e: e.update(handoff_id="wrong-handoff"))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(validator_operation="wrong-operation"))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(run_id="wrong-run"))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(predecessor_event_id="wrong-predecessor"))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(artifact_sha256="sha256:"+"0"*64))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(producer_agent="wrong-producer"))
        rejected(lambda s, e: s["validation_receipts"][e["receipt_ids"][0]].update(validator_agent="wrong-validator"))

def qa_defects_case():
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    from inspect_deck import render_previews, render_shape_masks, rendered_geometry_checks
    office=shutil.which("soffice") or shutil.which("libreoffice") or (r"C:\Program Files\LibreOffice\program\soffice.exe" if Path(r"C:\Program Files\LibreOffice\program\soffice.exe").is_file() else None)
    with tempfile.TemporaryDirectory() as td:
        td=Path(td); prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=RGBColor(255,255,255)
        overflow=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(1.2),Inches(.25)); overflow.name="overflow"
        overflow.text_frame.text="This deliberately overflows a tiny rendered text box by several complete lines"
        overflow.text_frame.paragraphs[0].runs[0].font.size=Pt(28)
        low=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1),Inches(4),Inches(3),Inches(.7)); low.name="low-contrast"
        low.fill.solid(); low.fill.fore_color.rgb=RGBColor(245,245,245); low.line.fill.background(); low.text="Low contrast"
        low.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(230,230,230)
        first=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(3),Inches(2),Inches(3),Inches(1)); first.name="overlap-a"; first.fill.solid(); first.fill.fore_color.rgb=RGBColor(220,20,60)
        second=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(4),Inches(2.2),Inches(3),Inches(1)); second.name="overlap-b"; second.fill.solid(); second.fill.fore_color.rgb=RGBColor(30,144,255)
        source=td/"source.png"
        source_image=Image.new("RGB",(100,100))
        source_pixels=source_image.load()
        bands=[(220,20,60),(255,140,0),(255,215,0),(34,139,34),(0,191,255),(30,144,255),(75,0,130),(148,0,211),(255,105,180),(10,30,60)]
        for x in range(100):
            for y in range(100):
                source_pixels[x,y]=bands[min(9,x//10)]
        source_image.save(source)
        pic=slide.shapes.add_picture(str(source),Inches(8),Inches(2),Inches(2),Inches(2)); pic.name="destructive-crop"; pic.crop_left=.95
        pptx=td/"defects.pptx"; prs.save(pptx)
        supported = bool(office and (shutil.which("pdftoppm") or importlib.util.find_spec("fitz") is not None))
        if supported:
            previews,error=render_previews(pptx,td/"previews")
            assert previews and not error
            masks=render_shape_masks(pptx,previews,td/"masks")
            mode="real-150dpi"
        else:
            # Controlled offline renderer fixture: exact 150-DPI canvas and
            # pixel masks model the same geometry without pretending that the
            # product rasterizer ran.
            from PIL import ImageDraw
            preview_path=td/"controlled-150dpi.png"
            canvas=Image.new("RGB",(2000,1125),(255,255,255)); draw=ImageDraw.Draw(canvas)
            draw.rectangle((150,150,520,260),fill=(20,20,20))  # visible overflow
            draw.rectangle((150,600,600,705),fill=(245,245,245)); draw.text((170,620),"Low contrast",fill=(230,230,230))
            draw.rectangle((450,300,900,450),fill=(220,20,60)); draw.rectangle((600,330,1050,480),fill=(30,144,255))
            draw.rectangle((1200,300,1500,600),fill=bands[-1])
            canvas.save(preview_path)
            previews=[{"slide_index":1,"path":str(preview_path),"sha256":sha_file(preview_path),"dpi":150}]
            masks={}
            for name,rect in {
                "overflow":(150,150,520,260), "low-contrast":(150,600,600,705),
                "overlap-a":(450,300,900,450), "overlap-b":(600,330,1050,480),
                "destructive-crop":(1200,300,1500,600),
            }.items():
                mask=Image.new("L",canvas.size,0); ImageDraw.Draw(mask).rectangle(rect,fill=255); masks[(1,name)]=mask
            mode="controlled-150dpi"
        token=load(PACK/"design-systems/registry.json")["systems"][0]
        checks,findings=rendered_geometry_checks(previews,pptx,token,masks)
        ids={x["id"] for x in findings}
        assert any(x.startswith("text-overflow-") for x in ids)
        assert any(x.startswith("contrast-") for x in ids), (ids, [x for x in checks if "contrast" in x[0]])
        assert any(x.startswith("crop-") for x in ids)
        assert any(x.startswith("rendered-overlap-") for x in ids)
        assert mode in {"real-150dpi","controlled-150dpi"}

def qa_case():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td); spec = td / "deck.json"; dump(spec, deck())
        run([sys.executable, str(SCRIPTS / "render_deck.py"), "--spec", str(spec), "--output-dir", str(td / "out")])
        model = td / "model.json"; dump(model, {"status": "completed", "scores": {"narrative": .9, "visual_craft": .9, "evidence_fit": .9}, "notes": ["Inspected hierarchy, composition, and evidence fit on every available preview."]})
        report = td / "qa.json"; receipts = td / "receipts"
        result = subprocess.run([sys.executable, str(SCRIPTS / "inspect_deck.py"), "--spec", str(spec), "--manifest", str(td / "out/render-manifest.json"), "--output", str(report), "--run-id", "run-fixture", "--handoff-id", "qa-h", "--receipt-dir", str(receipts), "--model-review", str(model)], cwd=REPO, capture_output=True, text=True)
        if result.returncode not in (0, 1):
            raise AssertionError(result.stdout + result.stderr)
        qa = load(report)
        assert not schema_validate(qa, load(PACK / "schemas/v3/qa-report.schema.json"))
        assert qa["model_review"]["status"] == "completed"
        if qa["deterministic_checks"][-2]["id"] == "visual-render-150dpi" and qa["deterministic_checks"][-2]["status"] == "unavailable":
            assert qa["verdict"] == "unverified-needs-user"
            assert qa["channels"]["pptx"] == "unavailable"
        emitted = [load(x) for x in receipts.glob("*.json")]
        assert len(emitted) == 4 and all(not schema_validate(x, load(PACK / "schemas/v3/validation-receipt.schema.json")) for x in emitted)

def state_case():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td); sp = td / "state.json"; ep = td / "event.json"; dump(sp, state())
        seq = 0
        def apply(typ, evidence=None, expected=0, **extra):
            nonlocal seq
            if extra.get("receipt_ids") and "handoff_id" not in extra:
                registry=load(sp)["validation_receipts"]
                extra["handoff_id"]=registry[extra["receipt_ids"][0]]["handoff_id"]
            seq += 1; event = {"event_id": f"e-{seq}-{typ}", "type": typ, "run_id": "run-fixture", "at": f"2026-01-01T00:00:{seq:02d}Z", "actor": "story-orchestrator", "evidence_ids": evidence or [], **extra}; dump(ep, event)
            run([sys.executable, str(SCRIPTS / "validate_contract.py"), "--state-transition", "--state", str(sp), "--event", str(ep), "--schema", str(PACK / "schemas/v3/state.schema.json")], expected)
            return event
        apply("initialize")
        assert load(sp)["phase"] == "initialized"
        current = load(sp); current["input_manifest"] = [{"artifact_id": "intake", "kind": "intake", "path": "intake.json", "sha256": "sha256:" + "a" * 64, "producer": "story-orchestrator", "status": "validated", "lineage_id": "lineage-1"}]; dump(sp, current)
        apply("validate-intake-started")
        current=load(sp); predecessor=current["events"][-1]["event_id"]
        current["validation_receipts"]["r-intake"]=indexed_receipt("r-intake","h-intake","intake","intake","sha256:"+"a"*64,"validate-intake",predecessor,producer="story-orchestrator",validator="deck-composer")
        current["validation_receipt_ids"]=["r-intake"]; dump(sp,current)
        apply("intake-validation-succeeded", ["r-intake","intake","sha256:"+"a"*64], receipt_ids=["r-intake"])
        apply("strategy-produced", ["r-intake"])
        current=load(sp); predecessor=current["events"][-1]["event_id"]
        strategy = [
            ("r-ledger","evidence-ledger","evidence-ledger","sha256:"+"4"*64),
            ("r-plan","story-plan","story-plan","sha256:"+"5"*64),
            ("r-proposal","proposal","proposal","sha256:"+"7"*64),
        ]
        current["artifact_manifest"]=[
            {"artifact_id":artifact_id,"kind":kind,"path":artifact_id+".json","sha256":checksum,
             "producer":"narrative-strategist","status":"validated","lineage_id":"lineage-1"}
            for _,artifact_id,kind,checksum in strategy
        ]
        for receipt_id,artifact_id,kind,checksum in strategy:
            current["validation_receipts"][receipt_id]=indexed_receipt(
                receipt_id,"h-strategy",kind,artifact_id,checksum,"validate-strategy",predecessor,
                producer="narrative-strategist",validator="deck-composer")
            current["validation_receipt_ids"].append(receipt_id)
        dump(sp,current)
        strategy_receipts=[x[0] for x in strategy]
        strategy_evidence=strategy_receipts+[x[1] for x in strategy]+[x[3] for x in strategy]
        apply("strategy-validation-succeeded", strategy_evidence, receipt_ids=strategy_receipts)
        apply("proposal-presented", strategy_receipts)
        approval = apply("approve", strategy_receipts, receipt_ids=strategy_receipts)
        apply("compose-started", [approval["event_id"],*strategy_receipts], approval_event_id=approval["event_id"], receipt_ids=strategy_receipts)
        current=load(sp); current["staging_manifest"]=[
            {"artifact_id":"deck-spec","kind":"deck-spec","path":"deck-spec.json","sha256":"sha256:"+"e"*64,"producer":"deck-composer","status":"staged","lineage_id":"lineage-1"},
            {"artifact_id":"deck-out","kind":"pptx","path":"out.pptx","sha256":"sha256:"+"d"*64,"producer":"deck-composer","status":"staged","lineage_id":"lineage-1"},
            {"artifact_id":"render-manifest","kind":"render-manifest","path":"render-manifest.json","sha256":"sha256:"+"b"*64,"producer":"deck-composer","status":"staged","lineage_id":"lineage-1"}]
        compose_event=next(x for x in reversed(current["events"]) if x["type"]=="compose-started")
        render_items=[
            ("r-spec","deck-spec","deck-spec","sha256:"+"e"*64),
            ("r-deck","deck-out","pptx","sha256:"+"d"*64),
            ("r-manifest","render-manifest","render-manifest","sha256:"+"b"*64),
        ]
        for receipt_id,artifact_id,kind,checksum in render_items:
            current["validation_receipts"][receipt_id]=indexed_receipt(
                receipt_id,"h-compose",kind,artifact_id,checksum,"compose",compose_event["event_id"],
                manifest="sha256:"+"b"*64,approval=approval["event_id"])
            current["validation_receipt_ids"].append(receipt_id)
        dump(sp,current)
        render_receipts=[x[0] for x in render_items]
        staged_ids=["deck-spec","deck-out","render-manifest"]
        apply("render-completed", [approval["event_id"],"cap-ok",*staged_ids,*render_receipts], approval_event_id=approval["event_id"], capability_evidence_id="cap-ok", receipt_ids=render_receipts, render_manifest_sha256="sha256:"+"b"*64)
        apply("review-started", [*staged_ids,*render_receipts], receipt_ids=render_receipts)
        current=load(sp); review_event=current["events"][-1]
        current["artifact_manifest"].append({"artifact_id":"qa-report","kind":"qa-report","path":"qa.json","sha256":"sha256:"+"6"*64,"producer":"deck-critic","status":"validated","lineage_id":"lineage-1"})
        current["validation_receipts"]["r-qa"]=indexed_receipt("r-qa","h-qa","qa-report","qa-report","sha256:"+"6"*64,"review",review_event["event_id"],producer="deck-critic",validator="deck-composer",capability="raster-unavailable",manifest="sha256:"+"b"*64)
        current["validation_receipt_ids"].append("r-qa"); dump(sp,current)
        # Complete transition-lineage negatives: every envelope dimension is
        # independently corrupted and must reject without advancing phase.
        baseline=sp.read_bytes()
        for field,value in [
            ("artifact_kind","story-plan"), ("artifact_id","story-plan"),
            ("artifact_sha256","sha256:"+"0"*64), ("validator_operation","compose"),
            ("producer_agent","narrative-strategist"), ("validator_agent","deck-critic-other"),
            ("run_id","other-run"), ("predecessor_event_id","unrelated-event"),
            ("lineage_id","stale-lineage"), ("capability_status","complete"),
            ("render_manifest_sha256","sha256:"+"0"*64),
        ]:
            broken=load(sp); broken["validation_receipts"]["r-qa"][field]=value; dump(sp,broken)
            apply("qa-unverified", [*staged_ids,"r-qa","sha256:"+"b"*64], expected=1, receipt_ids=["r-qa"], handoff_id="h-qa", render_manifest_sha256="sha256:"+"b"*64)
            assert load(sp)["phase"]=="reviewing"
            sp.write_bytes(baseline)
        apply("qa-unverified", [*staged_ids,"r-qa","sha256:"+"b"*64], expected=1, receipt_ids=["r-qa"], handoff_id="wrong-handoff", render_manifest_sha256="sha256:"+"b"*64)
        assert load(sp)["phase"]=="reviewing"; sp.write_bytes(baseline)
        missing=load(sp); del missing["validation_receipts"]["r-qa"]; dump(sp,missing)
        apply("qa-unverified", [*staged_ids,"r-qa","sha256:"+"b"*64], expected=1, receipt_ids=["r-qa"], handoff_id="h-qa", render_manifest_sha256="sha256:"+"b"*64)
        assert load(sp)["phase"]=="reviewing"; sp.write_bytes(baseline)
        apply("qa-unverified", ["deck-out","render-manifest","r-deck","sha256:"+"b"*64], expected=1, receipt_ids=["r-deck"], render_manifest_sha256="sha256:" + "b" * 64)
        qa = apply("qa-unverified", [*staged_ids,"r-qa","sha256:"+"b"*64], receipt_ids=["r-qa"], render_manifest_sha256="sha256:" + "b" * 64)
        apply("accept-residuals", ["visual-tooling",qa["event_id"],"r-qa","sha256:"+"c"*64], expected=1, acceptance_event_id=qa["event_id"], receipt_ids=["r-qa"], accepted_finding_ids=["visual-tooling"], render_manifest_sha256="sha256:" + "c" * 64)
        accepted = apply("accept-residuals", ["visual-tooling",qa["event_id"],"r-qa","sha256:"+"b"*64], acceptance_event_id=qa["event_id"], receipt_ids=["r-qa"], accepted_finding_ids=["visual-tooling"], render_manifest_sha256="sha256:" + "b" * 64)
        preflight_evidence=[accepted["event_id"],"sha256:"+"b"*64,"initial-destination",load(sp)["publication_destination"]["intake_sha256"],"r-qa"]
        apply("preflight-clear-residual", preflight_evidence, expected=1, receipt_ids=["r-qa"], acceptance_event_id="wrong", render_manifest_sha256="sha256:"+"b"*64, intake_sha256=load(sp)["publication_destination"]["intake_sha256"], preflight_fingerprint="sha256:"+"f"*64)
        apply("preflight-clear-residual", preflight_evidence, receipt_ids=["r-qa"], acceptance_event_id=accepted["event_id"], render_manifest_sha256="sha256:"+"b"*64, intake_sha256=load(sp)["publication_destination"]["intake_sha256"], preflight_fingerprint="sha256:"+"f"*64)
        apply("publication-committed", [accepted["event_id"],"r-deck","sha256:"+"b"*64], expected=1, receipt_ids=["r-deck"], acceptance_event_id=accepted["event_id"], render_manifest_sha256="sha256:" + "b" * 64)
        apply("publication-committed", [accepted["event_id"],"r-qa","sha256:"+"b"*64], receipt_ids=["r-qa"], acceptance_event_id=accepted["event_id"], render_manifest_sha256="sha256:" + "b" * 64)
        assert load(sp)["phase"] == "delivered-with-accepted-residuals"
        degraded=state("awaiting-render-decision"); degraded["approval_event_id"]="approved-e"; degraded["events"]=[{"event_id":"approved-e","type":"approve","run_id":"run-fixture","at":"2026-01-01T00:00:00Z","actor":"story-orchestrator","evidence_ids":["r-strategy"],"receipt_ids":["r-strategy"],"handoff_id":"h-strategy"}]; degraded["artifact_manifest"]=[{"artifact_id":"story-plan","kind":"story-plan","path":"story-plan.json","sha256":"sha256:"+"5"*64,"producer":"narrative-strategist","status":"validated","lineage_id":"lineage-1"}]; degraded["validation_receipt_ids"]=["r-strategy"]; degraded["validation_receipts"]["r-strategy"]=indexed_receipt("r-strategy","h-strategy","story-plan","story-plan","sha256:"+"5"*64,"validate-strategy","strategy-produced",producer="narrative-strategist",validator="deck-composer",approval="approved-e"); degraded["capability_preflight"].update({"status":"degraded","evidence_id":"cap-degraded","available_formats":"marp"}); degraded["staging_manifest"]=[{"artifact_id":"failed-pptx","kind":"pptx","path":"out.pptx","sha256":"sha256:"+"9"*64,"producer":"deck-composer","status":"staged","lineage_id":"lineage-1"}]; dump(sp,degraded)
        before=sp.read_bytes()
        apply("accept-degraded-format", ["cap-degraded","approved-e"], expected=1, capability_evidence_id="cap-degraded", effective_formats="pptx", invalidate_artifact_ids=["failed-pptx"])
        assert load(sp)["effective_formats"]=="both"
        apply("accept-degraded-format", ["cap-degraded","approved-e"], capability_evidence_id="cap-degraded", effective_formats="marp", invalidate_artifact_ids=["failed-pptx"])
        assert load(sp)["effective_formats"]=="marp" and load(sp)["staging_manifest"][0]["status"]=="invalidated"
        stale=state("approved"); stale["approval_event_id"]="approved-e"; stale["artifact_manifest"]=[{"artifact_id":"story-plan","kind":"story-plan","path":"story-plan.json","sha256":"sha256:"+"5"*64,"producer":"narrative-strategist","status":"validated","lineage_id":"lineage-1"},{"artifact_id":"unrelated","kind":"story-plan","path":"other.json","sha256":"sha256:"+"7"*64,"producer":"narrative-strategist","status":"validated","lineage_id":"lineage-1"}]; stale["validation_receipt_ids"]=["r-strategy","r-unrelated"]; stale["validation_receipts"]={"r-strategy":indexed_receipt("r-strategy","h-strategy","story-plan","story-plan","sha256:"+"5"*64,"validate-strategy","strategy-produced",producer="narrative-strategist",validator="deck-composer",approval="approved-e"),"r-unrelated":indexed_receipt("r-unrelated","h-other","story-plan","unrelated","sha256:"+"7"*64,"validate-strategy","other-strategy",producer="narrative-strategist",validator="deck-composer")}; stale["events"]=[{"event_id":"approved-e","type":"approve","run_id":"run-fixture","at":"2026-01-01T00:00:00Z","actor":"story-orchestrator","evidence_ids":["r-strategy"],"receipt_ids":["r-strategy"],"handoff_id":"h-strategy"}]; dump(sp,stale)
        apply("compose-started", ["approved-e","r-unrelated"], expected=1, approval_event_id="approved-e", receipt_ids=["r-unrelated"])
        assert load(sp)["phase"]=="approved"
        resume_state = state("rendered"); resume_state["validation_receipt_ids"] = ["r-strategy"]; resume_state["artifact_manifest"] = [{"artifact_id": "stale", "kind": "deck", "path": "x", "sha256": "sha256:" + "c" * 64, "producer": "deck-composer", "status": "staged", "lineage_id": "lineage-1"}]; resume_state["validation_receipts"]["r-strategy"]=indexed_receipt("r-strategy","h-resume","deck","stale","sha256:"+"c"*64,"resume","prior-checkpoint"); dump(sp, resume_state)
        apply("resume", ["r-strategy"], receipt_ids=["r-strategy"], resume_phase="strategy-ready", invalidate_artifact_ids=["stale"])
        assert load(sp)["phase"] == "strategy-ready" and "stale" in load(sp)["invalidated_artifact_ids"]
        resume_state = state("reviewing"); resume_state["validation_receipt_ids"]=["r-deck"]; resume_state["staging_manifest"]=[{"artifact_id":"render-manifest","kind":"render-manifest","path":"m","sha256":"sha256:"+"b"*64,"producer":"deck-composer","status":"staged","lineage_id":"lineage-1"}]; resume_state["validation_receipts"]["r-deck"]=indexed_receipt("r-deck","h-resume","render-manifest","render-manifest","sha256:"+"b"*64,"resume","prior-checkpoint",manifest="sha256:"+"b"*64); dump(sp,resume_state)
        apply("resume", ["r-deck"], expected=1, receipt_ids=["r-deck"], resume_phase="rendered", invalidate_artifact_ids=["old"])
        apply("resume", ["r-deck","sha256:"+"b"*64], receipt_ids=["r-deck"], render_manifest_sha256="sha256:"+"b"*64, resume_phase="rendered", invalidate_artifact_ids=["old"])
        assert load(sp)["phase"]=="rendered"
        passed=state("passed"); passed["qa_acceptance_event_id"]="qa-pass"; passed["validation_receipt_ids"]=["r-qa"]; passed["staging_manifest"]=[{"artifact_id":"render-manifest","kind":"render-manifest","path":"m","sha256":"sha256:"+"b"*64,"producer":"deck-composer","status":"staged","lineage_id":"lineage-1"}]; passed["artifact_manifest"]=[{"artifact_id":"qa-report","kind":"qa-report","path":"qa","sha256":"sha256:"+"6"*64,"producer":"deck-critic","status":"validated","lineage_id":"lineage-1"}];         passed["validation_receipts"]["r-qa"]=indexed_receipt("r-qa","h-qa","qa-report","qa-report","sha256:"+"6"*64,"review","review-started",producer="deck-critic",validator="deck-composer",capability="complete",manifest="sha256:"+"b"*64,acceptance="qa-pass"); passed["events"]=[
            {"event_id":"review-started","type":"review-started","run_id":"run-fixture","at":"2026-01-01T00:00:00Z","actor":"story-orchestrator","evidence_ids":["render-manifest"]},
            {"event_id":"qa-pass","type":"qa-passed","run_id":"run-fixture","at":"2026-01-01T00:00:01Z","actor":"story-orchestrator","evidence_ids":["render-manifest","r-qa","sha256:"+"b"*64],"receipt_ids":["r-qa"],"handoff_id":"h-qa","render_manifest_sha256":"sha256:"+"b"*64}
        ]; dump(sp,passed)
        apply("publish-requested", ["qa-pass","r-qa","sha256:"+"c"*64], expected=1, acceptance_event_id="qa-pass", receipt_ids=["r-qa"], render_manifest_sha256="sha256:"+"c"*64)
        wrong_acceptance=load(sp); wrong_acceptance["validation_receipts"]["r-qa"]["acceptance_event_id"]="other-acceptance"; dump(sp,wrong_acceptance)
        apply("publish-requested", ["qa-pass","r-qa","sha256:"+"b"*64], expected=1, acceptance_event_id="qa-pass", receipt_ids=["r-qa"], render_manifest_sha256="sha256:"+"b"*64)
        dump(sp,passed)
        apply("publish-requested", ["qa-pass","r-qa","sha256:"+"b"*64], acceptance_event_id="qa-pass", receipt_ids=["r-qa"], render_manifest_sha256="sha256:"+"b"*64)
        assert load(sp)["phase"]=="publication-preflighting"
        collision = state("publication-preflighting"); collision["publication_gate"] = "pass"; collision["publication_manifest_sha256"]="sha256:"+"b"*64; collision["qa_acceptance_event_id"]="qa-pass"; collision["validation_receipt_ids"]=["r-qa"]; collision["input_manifest"] = [{"artifact_id": "intake", "kind": "intake", "path": "intake", "sha256": collision["publication_destination"]["intake_sha256"], "producer": "story-orchestrator", "status": "validated", "lineage_id": "lineage-1"}]; collision["artifact_manifest"]=[{"artifact_id":"qa-report","kind":"qa-report","path":"qa","sha256":"sha256:"+"6"*64,"producer":"deck-critic","status":"validated","lineage_id":"lineage-1"}];         collision["validation_receipts"]["r-qa"]=indexed_receipt("r-qa","h-qa","qa-report","qa-report","sha256:"+"6"*64,"review","review-started",producer="deck-critic",validator="deck-composer",capability="complete",manifest="sha256:"+"b"*64,acceptance="qa-pass"); collision["events"]=[
            {"event_id":"review-started","type":"review-started","run_id":"run-fixture","at":"2026-01-01T00:00:00Z","actor":"story-orchestrator","evidence_ids":["render-manifest"]},
            {"event_id":"qa-pass","type":"qa-passed","run_id":"run-fixture","at":"2026-01-01T00:00:01Z","actor":"story-orchestrator","evidence_ids":["r-qa","sha256:"+"b"*64],"receipt_ids":["r-qa"],"handoff_id":"h-qa","render_manifest_sha256":"sha256:"+"b"*64}
        ]; dump(sp, collision)
        lineage=["sha256:"+"b"*64,"qa-pass","initial-destination",collision["publication_destination"]["intake_sha256"],"r-qa"]
        collision_event=apply("preflight-collision", lineage, receipt_ids=["r-qa"], acceptance_event_id="qa-pass", render_manifest_sha256="sha256:"+"b"*64, intake_sha256=collision["publication_destination"]["intake_sha256"], preflight_fingerprint="sha256:"+"c"*64)
        apply("replace-authorized", [collision_event["event_id"],"sha256:"+"b"*64,collision["publication_destination"]["intake_sha256"]], expected=1, intake_sha256=collision["publication_destination"]["intake_sha256"], preflight_fingerprint="sha256:"+"d"*64, invalidate_artifact_ids=["old-preflight"])
        override = apply("replace-authorized", [collision_event["event_id"],"sha256:"+"b"*64,collision["publication_destination"]["intake_sha256"]], intake_sha256=collision["publication_destination"]["intake_sha256"], preflight_fingerprint="sha256:"+"c"*64, invalidate_artifact_ids=["old-preflight"])
        apply("preflight-clear", ["sha256:"+"b"*64,"qa-pass","initial-destination",collision["publication_destination"]["intake_sha256"],"r-qa"], receipt_ids=["r-qa"], acceptance_event_id="qa-pass", render_manifest_sha256="sha256:"+"b"*64, intake_sha256=collision["publication_destination"]["intake_sha256"], preflight_fingerprint="sha256:"+"e"*64)
        assert load(sp)["phase"] == "publishing"
        invalid = {"event_id": "e-bad", "type": "made-up", "run_id": "run-fixture", "at": "2026-01-01T00:00:02Z", "actor": "x", "evidence_ids": []}; dump(ep, invalid)
        run([sys.executable, str(SCRIPTS / "validate_contract.py"), "--state-transition", "--state", str(sp), "--event", str(ep), "--schema", str(PACK / "schemas/v3/state.schema.json")], 1)
        after = load(sp); assert after["phase"] == "publishing" and after["events"][-1]["type"] == "transition-rejected"

def legacy_case():
    with tempfile.TemporaryDirectory() as td:
        run([sys.executable, str(SCRIPTS / "install-legacy-layout.py"), "--target", td])
        run([sys.executable, str(SCRIPTS / "install-legacy-layout.py"), "--target", td, "--check"])
        target=Path(td)/".github/agents/story-orchestrator.agent.md"
        target.write_text("collision\n",encoding="utf-8")
        before=target.read_bytes()
        run([sys.executable, str(SCRIPTS / "install-legacy-layout.py"), "--target", td, "--check"],1)
        result=subprocess.run([sys.executable,str(SCRIPTS/"install-legacy-layout.py"),"--target",td],cwd=REPO,capture_output=True,text=True)
        assert result.returncode != 0 and target.read_bytes()==before
        run([sys.executable,str(SCRIPTS/"install-legacy-layout.py"),"--target",td,"--force"])
        run([sys.executable,str(SCRIPTS/"install-legacy-layout.py"),"--target",td,"--check"])

def publication_case():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td); staged = td / "staged"; dest = td / "dest"; staged.mkdir(); dest.mkdir()
        output = staged / "output.md"; output.write_text("# Published\n", encoding="utf-8")
        manifest = {"schema_version": "3.0", "deck_id": "d", "input_checksums": {"deck_spec": "sha256:" + "2" * 64}, "renderer": {"name": "fixture", "version": "1", "scripts": {}}, "environment": {"python": sys.version.split()[0], "platform": sys.platform, "dependencies": {}, "tools": {}, "fonts": {"requested": {}, "substitutions": []}}, "render_seed": 1, "design_registry_checksum": "sha256:" + "3" * 64, "outputs": [{"channel": "marp", "path": str(output.resolve()), "sha256": sha_file(output), "status": "rendered", "semantic_hash": "sha256:" + "4" * 64, "diagnostics": []}], "slide_count": 1, "dimensions": {"aspect_ratio": "16:9", "width_in": 13.333, "height_in": 7.5}, "warnings": [], "capabilities": {}, "parity": {"status": "not-applicable", "matching": None}, "reproducibility": {"model_semantic_hash": "sha256:" + "5" * 64, "normalized_properties": [], "environment_sensitive": [], "second_render_checked": True, "matching": True, "first_normalized_hashes": {"marp": sha_file(output)}, "second_normalized_hashes": {"marp": sha_file(output)}}, "previews": []}
        mp = staged / "render-manifest.json"; dump(mp, manifest); manifest_sha = sha_file(mp)
        qa = {"schema_version": "3.0", "run_id": "run-fixture", "handoff_id": "qa-h", "verdict": "pass", "channels": {"pptx": "not-requested", "marp": "pass", "parity": "not-applicable"}, "deterministic_checks": [{"id": "all", "status": "pass", "evidence": "fixture"}], "model_review": {"status": "completed", "scores": {"narrative": 1, "visual_craft": 1, "evidence_fit": 1}, "notes": ["inspected"]}, "findings": [], "metrics": {"structural": 1, "source_coverage": 1, "narrative": 1, "evidence_fit": 1, "accessibility": 1, "visual_craft": 1, "parity": 1, "reproducibility": 1}, "prior_comparison": None, "validated_receipt_ids": [], "previews": [], "reproducibility": {"checked": True, "matching": True, "comparison_manifest_sha256": None}}
        qp = td / "qa.json"; dump(qp, qa)
        acceptance = {"event_id": "qa-accepted", "type": "qa-passed", "run_id": "run-fixture", "at": "2026-01-01T00:00:02Z", "actor": "story-orchestrator", "evidence_ids": ["qa"], "render_manifest_sha256": manifest_sha}
        ap = td / "acceptance.json"; dump(ap, acceptance)
        destination_event = {"event_id": "dest-e", "type": "destination-selected", "run_id": "run-fixture", "at": "2026-01-01T00:00:01Z", "actor": "story-orchestrator", "evidence_ids": ["intake"], "destination": str(dest.resolve()), "intake_sha256": "sha256:" + "1" * 64, "invalidate_artifact_ids": ["old"]}
        dp = td / "destination.json"; dump(dp, destination_event)
        receipt_core = {"run_id": "run-fixture", "handoff_id": "qa-h", "artifact_kind": "qa-report", "artifact_path": str(qp.resolve()), "artifact_sha256": sha_file(qp), "schema_sha256": sha_file(PACK / "schemas/v3/qa-report.schema.json"), "producer_agent": "deck-critic", "validator_agent": "deck-composer", "validator_operation": "publish", "diagnostics_hash": sha_bytes(canonical([]))}
        receipt_core["inspection_profile"]=None
        receipt = {"schema_version": "3.0", "receipt_version": "1.0", "receipt_id": sha_bytes(canonical(receipt_core)), **receipt_core, "schema_path": str((PACK / "schemas/v3/qa-report.schema.json").resolve()), "tool_version": "fixture", "validated_at": "2026-01-01T00:00:03Z", "valid": True, "diagnostic_class": "none", "diagnostic_count": 0, "diagnostics": [], "superseded": False, "accepted_qa_event_id": "qa-accepted", "accepted_residual_event_id": None, "render_manifest_sha256": manifest_sha}
        rp = td / "qa-receipt.json"; dump(rp, receipt)
        acceptance["receipt_ids"]=[receipt["receipt_id"]]; acceptance["evidence_ids"]=[receipt["receipt_id"],manifest_sha]; dump(ap,acceptance)
        st = state("publishing"); st["effective_formats"] = "marp"; st["requested_formats"] = "marp"; st["publication_gate"] = "pass"; st["publication_manifest_sha256"] = manifest_sha; st["qa_acceptance_event_id"] = "qa-accepted"; st["publication_destination"].update({"effective_output_dir": str(dest.resolve()), "effective_destination_event_id": "dest-e"}); st["staging_manifest"] = [{"artifact_id": "out", "kind": "marp", "path": str(output.resolve()), "sha256": sha_file(output), "producer": "deck-composer", "status": "staged", "lineage_id": "lineage-1"}]; st["artifact_manifest"]=[{"artifact_id":"qa-report","kind":"qa-report","path":str(qp.resolve()),"sha256":sha_file(qp),"producer":"deck-critic","status":"validated","lineage_id":"lineage-1"}]; st["validation_receipt_ids"]=[receipt["receipt_id"]]; st["validation_receipts"][receipt["receipt_id"]]=indexed_receipt(receipt["receipt_id"],"qa-h","qa-report","qa-report",sha_file(qp),"publish","qa-accepted",producer="deck-critic",validator="deck-composer",capability="complete",manifest=manifest_sha,acceptance="qa-accepted"); st["events"] = [destination_event, acceptance]
        sp = td / "state.json"; dump(sp, st)
        pre = run([sys.executable, str(SCRIPTS / "publish_artifacts.py"), "--destination", str(dest), "--manifest", str(mp), "--preflight"])
        fp = json.loads(pre.stdout)["metadata_fingerprint"]
        command = [sys.executable, str(SCRIPTS / "publish_artifacts.py"), "--destination", str(dest), "--manifest", str(mp), "--expected-fingerprint", fp, "--qa-report", str(qp), "--qa-receipt", str(rp), "--acceptance-event", str(ap), "--destination-event", str(dp), "--intake-sha256", st["publication_destination"]["intake_sha256"], "--state", str(sp)]
        # Every receipt identity/diagnostic/handoff/run/acceptance/schema/manifest
        # mismatch must fail before destination mutation.
        for key,value in [
            ("receipt_id","sha256:"+"0"*64), ("diagnostics_hash","sha256:"+"0"*64),
            ("diagnostic_count",1), ("handoff_id","other"), ("run_id","other"),
            ("accepted_qa_event_id","other"), ("schema_sha256","sha256:"+"0"*64),
            ("render_manifest_sha256","sha256:"+"0"*64),
        ]:
            tampered=json.loads(json.dumps(receipt)); tampered[key]=value
            tp=td/f"tampered-{key}.json"; dump(tp,tampered)
            bad=list(command); bad[bad.index(str(rp))]=str(tp)
            before=list(dest.iterdir()); run(bad,8); assert list(dest.iterdir())==before
        run(command)
        current = dest / "current.json"; assert current.is_file()
        current.unlink()
        pre = run([sys.executable, str(SCRIPTS / "publish_artifacts.py"), "--destination", str(dest), "--manifest", str(mp), "--preflight"])
        command[command.index(fp)] = json.loads(pre.stdout)["metadata_fingerprint"]
        run(command); assert current.is_file()  # committed-existing repairs the pointer
        final=Path(load(current)["version_path"]); committed_manifest=final/"render-manifest.json"
        original_manifest=committed_manifest.read_bytes(); committed_manifest.write_text("{}\n",encoding="utf-8"); current.unlink()
        pre=run([sys.executable,str(SCRIPTS/"publish_artifacts.py"),"--destination",str(dest),"--manifest",str(mp),"--preflight"])
        command[command.index("--expected-fingerprint")+1]=json.loads(pre.stdout)["metadata_fingerprint"]
        run(command,9)
        committed_manifest.write_bytes(original_manifest)
        # Injected failures prove rollback after directory rename and after pointer replacement.
        if current.exists(): current.unlink()
        import shutil
        shutil.rmtree(final)
        pre=run([sys.executable,str(SCRIPTS/"publish_artifacts.py"),"--destination",str(dest),"--manifest",str(mp),"--preflight"])
        command[command.index("--expected-fingerprint")+1]=json.loads(pre.stdout)["metadata_fingerprint"]
        run(command+["--failpoint","post-rename"],10)
        assert not current.exists() and not final.exists()
        old_pointer=b'{"publication_id":"last-good"}\n'; current.write_bytes(old_pointer)
        pre=run([sys.executable,str(SCRIPTS/"publish_artifacts.py"),"--destination",str(dest),"--manifest",str(mp),"--preflight"])
        command[command.index("--expected-fingerprint")+1]=json.loads(pre.stdout)["metadata_fingerprint"]
        policy_command=command+["--policy","replace","--failpoint","pointer-write"]
        run(policy_command,10)
        assert current.read_bytes()==old_pointer and not final.exists()

CASES = {"contract": contract_case, "render": render_case, "render-variants": render_variants_case, "all-charts": all_charts_case, "all-recipes": all_recipes_case, "receipt": receipt_case, "receipt-diagnostics": receipt_diagnostics_case, "receipt-gates": receipt_gate_matrix_case, "qa": qa_case, "qa-defects": qa_defects_case, "state": state_case, "legacy": legacy_case, "publication": publication_case}
if __name__ == "__main__":
    args = argparse.ArgumentParser(); args.add_argument("case", choices=sorted(CASES)); selected = args.parse_args().case
    CASES[selected](); print(json.dumps({"case": selected, "status": "pass"}))
