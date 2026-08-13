"""Validate once, render sibling projections, and record output-derived semantics."""
import argparse, io, json, os, platform, re, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent))
from common import canonical, dump, file_version, libreoffice_cli, libreoffice_headless_command, load, quiet_subprocess_kwargs, semantic_slide, sha_bytes, sha_file

def extract_pptx(path):
 from pptx import Presentation
 out=[]
 for i,slide in enumerate(Presentation(path).slides,1):
  note=slide.notes_slide.notes_text_frame.text.strip()
  try: meta=json.loads(note[note.find("{"):])
  except Exception: meta={}
  title=next((x.text for x in slide.shapes if getattr(x,"name","")=="title"),"")
  out.append({"index":i,"slide_id":meta.get("slide_id"),"title":title,"claim_ids":meta.get("claim_ids",[]),"evidence_ids":meta.get("evidence_ids",[]),"source_footer":meta.get("source_footer",[]),"notes":meta.get("notes",{}),"cta":meta.get("cta")})
 return out
def extract_marp(path):
 return [json.loads(x) for x in re.findall(r"<!-- story-semantic: (\{.*\}) -->",Path(path).read_text(encoding="utf-8"))]
def command_version(cmd):
 try: return subprocess.run([cmd,"--version"],capture_output=True,text=True,timeout=5,**quiet_subprocess_kwargs()).stdout.strip().splitlines()[0]
 except Exception: return None
def normalized_output_hash(path,channel):
 if channel=="marp": return sha_file(path)
 def archive_rows(source):
  nested=[]
  with zipfile.ZipFile(source) as z:
   for name in sorted(z.namelist()):
    data=z.read(name)
    if name=="docProps/core.xml":
     text=data.decode("utf-8")
     text=re.sub(r"<dcterms:(created|modified)[^>]*>.*?</dcterms:\1>","",text)
     data=text.encode("utf-8")
    if name.endswith((".xlsx",".zip")):
     try: value=sha_bytes(canonical(archive_rows(io.BytesIO(data))))
     except zipfile.BadZipFile: value=sha_bytes(data)
    else: value=sha_bytes(data)
    nested.append((name,value))
  return nested
 rows=archive_rows(path)
 return sha_bytes(canonical(rows))
def font_available(name):
 normalized=re.sub(r"[^a-z0-9]","",name.casefold())
 roots=[Path(os.environ.get("WINDIR","C:/Windows"))/"Fonts",Path.home()/".fonts",Path("/usr/share/fonts")]
 if any(normalized in re.sub(r"[^a-z0-9]","",p.stem.casefold()) for root in roots if root.exists() for p in root.rglob("*") if p.is_file()):
  return True
 fc=shutil.which("fc-match")
 if fc:
  out=subprocess.run([fc,"--format=%{family}",name],capture_output=True,text=True).stdout.casefold()
  return name.casefold() in out
 return False
def make_previews(pptx,out):
 office=libreoffice_cli(); raster=shutil.which("pdftoppm")
 try: import fitz
 except ImportError: fitz=None
 if not office or (not raster and not fitz): return []
 out.mkdir(exist_ok=True)
 with tempfile.TemporaryDirectory() as td:
  profile=Path(td)/"profile"
  subprocess.run(libreoffice_headless_command(office,profile,"--convert-to","pdf","--outdir",td,str(Path(pptx).resolve())),check=True,capture_output=True,timeout=120,**quiet_subprocess_kwargs())
  pdf=Path(td)/(Path(pptx).stem+".pdf")
  if raster: subprocess.run([raster,"-png","-r","150",str(pdf),str(out/"slide")],check=True,capture_output=True)
  else:
   with fitz.open(pdf) as document:
    for index,page in enumerate(document,1): page.get_pixmap(matrix=fitz.Matrix(150/72,150/72),alpha=False).save(out/f"slide-{index}.png")
 return [{"slide_index":i,"path":str(p.resolve()),"sha256":sha_file(p),"dpi":150} for i,p in enumerate(sorted(out.glob("slide-*.png")),1)]
def main():
 p=argparse.ArgumentParser(); p.add_argument("--spec",required=True); p.add_argument("--output-dir",required=True); p.add_argument("--registry"); a=p.parse_args()
 spec=load(a.spec); root=Path(__file__).resolve().parents[1]; registry=Path(a.registry) if a.registry else root/"design-systems"/"registry.json"; out=Path(a.output_dir); out.mkdir(parents=True,exist_ok=True)
 # Validation is a hard prerequisite. A temporary receipt avoids trusting renderer-side ad hoc checks.
 validation=out/".render-validation-receipt.json"
 cmd=[sys.executable,str(Path(__file__).with_name("validate_contract.py")),"--kind","deck-spec","--input",str(Path(a.spec).resolve()),"--schema",str(root/"schemas/v3/deck-spec.schema.json"),"--run-id",spec["metadata"].get("created_from_run_id","render"),"--handoff-id","render-"+sha_file(a.spec)[7:23],"--producer","deck-composer","--validator","story-renderer","--operation","pre-render","--receipt",str(validation)]
 vr=subprocess.run(cmd,capture_output=True,text=True)
 if vr.returncode or not load(validation)["valid"]:
  print(vr.stdout or vr.stderr); return 2
 validation.unlink()
 requested=["pptx","marp"] if spec["formats"]=="both" else [spec["formats"]]; outputs=[]; diagnostics=[]; extracted={}
 office=libreoffice_cli()
 capabilities={"pptx":{"available":file_version("python-pptx") is not None},"marp":{"source_available":True,"cli_available":bool(shutil.which("marp") or shutil.which("marp-cli"))},"libreoffice":{"available":bool(office)},"rasterizer":{"available":bool(shutil.which("pdftoppm") or file_version("PyMuPDF"))}}
 for channel in requested:
  dest=out/("output.pptx" if channel=="pptx" else "output.md"); script=Path(__file__).with_name("render_"+channel+".py")
  if channel=="pptx" and not capabilities["pptx"]["available"]:
   outputs.append({"channel":channel,"path":str(dest.resolve()),"sha256":None,"status":"unavailable","semantic_hash":None,"diagnostics":["python-pptx unavailable"]}); continue
  run=subprocess.run([sys.executable,str(script),"--spec",str(Path(a.spec).resolve()),"--output",str(dest),"--registry",str(registry)],capture_output=True,text=True)
  if run.returncode:
   status="unavailable" if run.returncode==3 else "failed"; diag=(run.stdout+run.stderr).strip()
   outputs.append({"channel":channel,"path":str(dest.resolve()),"sha256":None,"status":status,"semantic_hash":None,"diagnostics":[diag]}); diagnostics.append(f"{channel}: {diag}"); continue
  try: sem=extract_pptx(dest) if channel=="pptx" else extract_marp(dest); extracted[channel]=sem
  except Exception as e:
   outputs.append({"channel":channel,"path":str(dest.resolve()),"sha256":sha_file(dest),"status":"failed","semantic_hash":None,"diagnostics":[f"semantic extraction failed: {e}"]}); continue
  outputs.append({"channel":channel,"path":str(dest.resolve()),"sha256":sha_file(dest),"status":"rendered","semantic_hash":sha_bytes(canonical(sem)),"diagnostics":[]})
 parity={"status":"not-applicable","matching":None}
 if {"pptx","marp"}.issubset(extracted):
  parity={"status":"pass" if extracted["pptx"]==extracted["marp"] else "fail","matching":extracted["pptx"]==extracted["marp"],"pptx_hash":sha_bytes(canonical(extracted["pptx"])),"marp_hash":sha_bytes(canonical(extracted["marp"]))}
 fonts=next(x["fonts"] for x in load(registry)["systems"] if x["id"]==spec["design_system_id"])
 substitutions=[]
 for role,stack in fonts.items():
  selected=next((name for name in stack if font_available(name)),None)
  if selected!=stack[0]: substitutions.append({"role":role,"requested":stack[0],"selected":selected,"reason":"requested font unavailable"})
 previews=[]
 if outputs and next((x for x in outputs if x["channel"]=="pptx" and x["status"]=="rendered"),None):
  try: previews=make_previews(out/"output.pptx",out/"previews-150dpi")
  except Exception as exc: diagnostics.append("preview generation failed: "+str(exc))
 normalized_first={x["channel"]:normalized_output_hash(x["path"],x["channel"]) for x in outputs if x["status"]=="rendered"}
 normalized_second={}; second_diagnostics=[]
 with tempfile.TemporaryDirectory() as td:
  for channel in requested:
   first=next((x for x in outputs if x["channel"]==channel and x["status"]=="rendered"),None)
   if not first: continue
   second=Path(td)/Path(first["path"]).name
   run=subprocess.run([sys.executable,str(Path(__file__).with_name("render_"+channel+".py")),"--spec",str(Path(a.spec).resolve()),"--output",str(second),"--registry",str(registry)],capture_output=True,text=True)
   if run.returncode: second_diagnostics.append(f"{channel} second render failed: {(run.stdout+run.stderr).strip()}")
   else: normalized_second[channel]=normalized_output_hash(second,channel)
 reproducible=not second_diagnostics and normalized_first==normalized_second
 manifest={"schema_version":"3.0","deck_id":spec["deck_id"],"input_checksums":{"deck_spec":sha_file(a.spec)},"renderer":{"name":"story-renderer","version":"3.2","scripts":{x.name:sha_file(x) for x in (Path(__file__),Path(__file__).with_name("render_pptx.py"),Path(__file__).with_name("render_marp.py"))}},"environment":{"python":platform.python_version(),"platform":platform.platform(),"dependencies":{"python-pptx":file_version("python-pptx"),"Pillow":file_version("Pillow"),"PyMuPDF":file_version("PyMuPDF"),"jsonschema":file_version("jsonschema")},"tools":{"marp":command_version(shutil.which("marp") or shutil.which("marp-cli")) if capabilities["marp"]["cli_available"] else None,"libreoffice":command_version(office) if capabilities["libreoffice"]["available"] else None}, "fonts":{"requested":fonts,"substitutions":substitutions}},"render_seed":spec["render_seed"],"design_registry_checksum":sha_file(registry),"outputs":outputs,"slide_count":len(spec["slides"]),"dimensions":{"aspect_ratio":"16:9","width_in":13.333,"height_in":7.5},"warnings":diagnostics+second_diagnostics,"capabilities":capabilities,"parity":parity,"reproducibility":{"model_semantic_hash":sha_bytes(canonical([semantic_slide(x) for x in spec["slides"]])),"normalized_properties":["PPTX ZIP member bytes excluding core created/modified timestamps","Marp bytes","object ordering","dimensions","design registry","seed"],"environment_sensitive":["font substitution","LibreOffice rasterization"],"second_render_checked":True,"matching":reproducible,"first_normalized_hashes":normalized_first,"second_normalized_hashes":normalized_second},"previews":previews}
 dump(out/"render-manifest.json",manifest); print(json.dumps(manifest,indent=2)); return 0 if all(x["status"]=="rendered" for x in outputs) and parity["status"]!="fail" and reproducible else 4
if __name__=="__main__": raise SystemExit(main())
